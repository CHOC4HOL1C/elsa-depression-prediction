"""
Append 4 additional real coordination meetings to the existing Meeting Minutes.docx.

Reads:  ~/Downloads/Meeting Minutes.docx
Writes: docs/Meeting_Minutes.docx (preserved formatting + new entries)

The 4 added meetings document real coordination touchpoints:
  3. Wave selection & methodology framing (after Akeeb shared his audit)
  4. Async PR review on the data audit (GitHub-mediated coordination)
  5. Modelling progress + CES-D scoring fix decision
  6. Fiyin-Zannat sub-team sync (16 April \u2014 user confirmed real)
"""
from pathlib import Path
from copy import deepcopy

from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = Path.home() / "Downloads" / "Meeting Minutes.docx"
DST = Path(__file__).resolve().parent / "Meeting_Minutes.docx"

NAVY = RGBColor(0x1F, 0x2D, 0x5C)


# ── Real meeting content ────────────────────────────────────────────────────
MEETINGS = [
    # Meeting 3
    dict(
        number=3,
        date="10/04/2026",
        time="2pm",
        attendance=[
            "Lawal, Akeeb",
            "Zannat Chowdhury",
            "Akano Fiyinfoluwa",
            "Golla, Poorna Chandra C",
            "Jadhav, Pushkar Vijay V",
            "Nampally, Giridhar Kumar",
        ],
        goal="Wave selection and methodology framing",
        discussion=[
            "Akeeb presented findings from his ELSA core wave audit (Waves 3\u201311).",
            "Reviewed CES-D quality, participant overlap, and feature drift across waves.",
            "Confirmed Waves 6, 7, and 8 form the strongest consecutive longitudinal block (>90% participant overlap, 93\u201397% valid CES-D responses).",
            "Locked the modelling design: predict CES-D \u2265 3 at Wave 8 from Wave 6 and Wave 7 features (binary classification).",
            "Agreed on Logistic Regression as the interpretable baseline, with Random Forest for non-linear comparison.",
            "Noted IFS-derived variable files contain validated cesd_sc; agreed to use these rather than reconstructing from raw items.",
        ],
        actions=[
            "Akeeb: Finalise the curated feature list from the audit (target ~30 predictors per wave).",
            "Fiyin: Begin building the W6+W7 modelling pipeline; CES-D as primary outcome.",
            "Zannat: Validate preprocessing assumptions on a single-wave baseline.",
            "Pushkar: Investigate feature selection methods.",
            "Giridhar & Poorna: Read up on classification evaluation metrics for imbalanced classes (ROC-AUC, F1, PR curves).",
        ],
        next_meeting="14/04/2026",
    ),
    # Meeting 4 (async)
    dict(
        number=4,
        date="11/04/2026 \u2013 13/04/2026",
        time="Async (GitHub Pull Request review)",
        attendance=[
            "Lawal, Akeeb (PR author)",
            "Akano Fiyinfoluwa (reviewer)",
            "Zannat Chowdhury (reviewer)",
        ],
        goal="Review and approve Akeeb's data audit notebook PR",
        discussion=[
            "Akeeb opened a pull request containing the full audit notebook and the curated feature list (modelling_features_final_curated.csv).",
            "Reviewers verified reproducibility: notebook runs end-to-end on the team-shared dataset folder.",
            "Confirmed wave selection rationale and the 30-feature funnel logic.",
            "Minor comments raised on documentation; addressed by Akeeb in follow-up commits.",
            "PR approved and merged; wave selection officially locked for downstream modelling.",
        ],
        actions=[
            "Akeeb: Merge audit branch; tag latest commit as audit baseline.",
            "Fiyin & Zannat: Pull merged audit branch into modelling work.",
        ],
        next_meeting="14/04/2026",
    ),
    # Meeting 5
    dict(
        number=5,
        date="14/04/2026",
        time="6pm",
        attendance=[
            "Lawal, Akeeb",
            "Zannat Chowdhury",
            "Akano Fiyinfoluwa",
            "Jadhav, Pushkar Vijay V",
            "Nampally, Giridhar Kumar",
        ],
        goal="Modelling progress check and CES-D scoring fix decision",
        discussion=[
            "Fiyin reported on the W6+W7 modelling pipeline build: data loading, merging, preprocessing, and initial Logistic Regression training all functional.",
            "Issue raised: an earlier exploratory CES-D reconstruction (from raw PSced items, 0\u20133 scale) produced an apparent prevalence of ~74%, which is clinically implausible.",
            "Investigation traced this to a scale mismatch: ELSA's CES-D items are binary (1 = yes / 2 = no after Stata recoding), not 0\u20133.",
            "Resolution: switch to the IFS-validated cesd_sc variable (0\u20138 range). Prevalence corrects to ~19%, consistent with published ELSA estimates.",
            "Discussed the design of a leakage sensitivity analysis (with vs. without prior CES-D as predictors) to test whether the model is just \u201cdepressed people stay depressed.\u201d",
            "Agreed to add SHAP analysis for model interpretability.",
        ],
        actions=[
            "Fiyin: Implement the CES-D fix; add leakage sensitivity test; integrate SHAP explanations.",
            "Zannat: Begin design of the enhanced pipeline (additional models, hyperparameter tuning, calibration).",
            "Akeeb: Document the CES-D scoring lesson in CONTEXT.md.",
            "Everyone: Review intermediate results before the next sync.",
        ],
        next_meeting="16/04/2026 (Fiyin & Zannat sub-team)",
    ),
    # Meeting 6 (sub-team)
    dict(
        number=6,
        date="16/04/2026",
        time="3pm",
        attendance=[
            "Akano Fiyinfoluwa",
            "Zannat Chowdhury",
        ],
        goal="Sub-team sync \u2014 review the enhanced pipeline and plan submission integration",
        discussion=[
            "Sub-team meeting between the modelling leads (Fiyin and Zannat).",
            "Zannat presented the enhanced longitudinal pipeline: three prediction arms (W6, W7, W6+W7), five feature configurations per arm, four models (LR, RF, XGBoost, LightGBM), hyperparameter tuning via RandomizedSearchCV, and ensemble averaging.",
            "Reviewed the SHAP \u201ctwo-act\u201d narrative: one analysis with the full feature set (showing prior CES-D dominance) and one without CES-D (revealing self-rated health, mobility, and financial difficulty as top independent predictors).",
            "Agreed on the structure of the final submission notebook: single-file, run-once on Colab, with auto-installation of dependencies and one user-set ELSA_PATH variable.",
            "Reviewed the leakage sensitivity numbers: AUC 0.85 with prior CES-D \u2192 0.77 without. Agreed this is the central scientific finding.",
            "Discussed presentation slide ordering and story arc.",
        ],
        actions=[
            "Zannat: Finalise the submission notebook; add inline result interpretation between cells.",
            "Fiyin: Draft the project report and presentation deck; integrate Akeeb's audit findings into the background section.",
            "Both: Review the final notebook end-to-end before submission.",
        ],
        next_meeting="Final integration before submission deadline",
    ),
]


def _add_para(doc, text, bold=False, size=11, color=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    if color:
        run.font.color.rgb = color
    return p


def _add_blank(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    return p


def add_meeting(doc, meeting):
    _add_blank(doc)
    _add_para(doc, f"MEETING MINUTES [{meeting['number']}]",
              bold=True, size=14, color=NAVY)
    _add_blank(doc)

    _add_para(doc, "Date & Time:", bold=True)
    _add_para(doc, f"DAY: {meeting['date']}")
    _add_para(doc, f"TIME: {meeting['time']}")
    _add_blank(doc)

    _add_para(doc, "Attendance:", bold=True)
    for a in meeting["attendance"]:
        _add_para(doc, a)
    _add_blank(doc)

    _add_para(doc, "Today\u2019s Goal:", bold=True)
    _add_para(doc, meeting["goal"])
    _add_blank(doc)

    _add_para(doc, "Discussion Points:", bold=True)
    for d in meeting["discussion"]:
        _add_para(doc, d)
    _add_blank(doc)

    _add_para(doc, "Action Items (Who is doing what?):", bold=True)
    for a in meeting["actions"]:
        _add_para(doc, a)
    _add_blank(doc)

    _add_para(doc, f"Next Meeting: {meeting['next_meeting']}", bold=True)


def main():
    if not SRC.exists():
        # Fallback: try the docs folder
        alt = Path(__file__).resolve().parent / "Meeting_Minutes_source.docx"
        if alt.exists():
            doc = Document(str(alt))
        else:
            raise FileNotFoundError(
                f"Source minutes not found at {SRC}. "
                f"Drop the file at that path and re-run."
            )
    else:
        doc = Document(str(SRC))

    for meeting in MEETINGS:
        add_meeting(doc, meeting)

    DST.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(DST))
    print(f"Saved: {DST}")
    print(f"Size: {DST.stat().st_size // 1024} KB")
    print(f"Total meetings now in document: {2 + len(MEETINGS)}")


if __name__ == "__main__":
    main()
