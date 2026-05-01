"""
Convert docs/REPORT.md into a polished Word document for sharing.

Usage:  python3 docs/build_report_docx.py
Output: docs/ELSA_Depression_Prediction_Report.docx

The script does NOT use pandoc. It performs a focused Markdown\u2192DOCX conversion
covering the constructs actually used in REPORT.md: headings (#, ##, ###),
paragraphs with bold/italic/inline-code, bulleted lists, fenced code blocks
(rendered as monospaced indented blocks), pipe tables, and horizontal rules.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SRC  = ROOT / "docs" / "REPORT.md"
OUT  = ROOT / "docs" / "ELSA_Depression_Prediction_Report.docx"

# ── Style palette ────────────────────────────────────────────────────────────
NAVY  = RGBColor(0x1F, 0x2D, 0x5C)
TEAL  = RGBColor(0x18, 0x6F, 0x82)
CORAL = RGBColor(0xE5, 0x6B, 0x5C)
GREY  = RGBColor(0x55, 0x5C, 0x66)
INK   = RGBColor(0x22, 0x26, 0x32)


def _set_cell_bg(cell, color_hex):
    """Set a table cell's shading colour (hex without #)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tc_pr.append(shd)


def _add_inline(paragraph, text, base_size=11, base_color=INK):
    """Render inline markdown: **bold**, *italic*, `code`."""
    pattern = re.compile(r"(\*\*[^\*]+\*\*|\*[^\*]+\*|`[^`]+`)")
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            run = paragraph.add_run(text[pos:m.start()])
            run.font.name = "Calibri"
            run.font.size = Pt(base_size)
            run.font.color.rgb = base_color
        token = m.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            run.font.name = "Calibri"
            run.font.size = Pt(base_size)
            run.font.color.rgb = base_color
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(base_size - 1)
            run.font.color.rgb = NAVY
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
            run.font.name = "Calibri"
            run.font.size = Pt(base_size)
            run.font.color.rgb = base_color
        pos = m.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        run.font.name = "Calibri"
        run.font.size = Pt(base_size)
        run.font.color.rgb = base_color


def _heading(doc, text, level):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    sizes = {1: 22, 2: 16, 3: 13}
    colors = {1: NAVY, 2: NAVY, 3: TEAL}
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(sizes.get(level, 12))
    run.font.color.rgb = colors.get(level, INK)
    if level == 1:
        # Add subtle underline bar by inserting a bottom border
        pPr = p._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "8")
        bottom.set(qn("w:space"), "4")
        bottom.set(qn("w:color"), "1F2D5C")
        bdr.append(bottom)
        pPr.append(bdr)
    return p


def _paragraph(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    _add_inline(p, text, base_size=11, base_color=INK)
    return p


def _bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.2
    _add_inline(p, text, base_size=11, base_color=INK)
    return p


def _code_block(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Cm(0.5)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F4F5F7")
    pPr.append(shd)
    for i, line in enumerate(lines):
        if i:
            p.add_run("\n")
        run = p.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(9.5)
        run.font.color.rgb = NAVY


def _table(doc, header, rows):
    n_cols = len(header)
    table = doc.add_table(rows=1 + len(rows), cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Light Grid Accent 1"
    table.autofit = True

    # Header row
    hdr_cells = table.rows[0].cells
    for j, h in enumerate(header):
        cell = hdr_cells[j]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        _set_cell_bg(cell, "1F2D5C")
        cell.paragraphs[0].text = ""
        run = cell.paragraphs[0].add_run(h.strip())
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size = Pt(10.5)
        run.font.name = "Calibri"

    # Data rows
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = table.rows[i].cells[j]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell.paragraphs[0].text = ""
            _add_inline(cell.paragraphs[0], val.strip(), base_size=10, base_color=INK)
    return table


def _hr(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "AAAAAA")
    bdr.append(bottom)
    pPr.append(bdr)


# ── Markdown parser (purpose-built for REPORT.md) ───────────────────────────
def parse_markdown(md: str):
    """Yield ('heading', level, text) | ('para', text) | ('bullet', text)
       | ('code', lines) | ('table', header, rows) | ('hr',) tokens."""
    lines = md.splitlines()
    i = 0
    n = len(lines)

    while i < n:
        line = lines[i]

        # Skip stray blank lines
        if not line.strip():
            i += 1
            continue

        # Horizontal rule
        if re.fullmatch(r"-{3,}", line.strip()):
            yield ("hr",)
            i += 1
            continue

        # Heading
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            yield ("heading", level, m.group(2).strip())
            i += 1
            continue

        # Fenced code block
        if line.strip().startswith("```"):
            body = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                body.append(lines[i])
                i += 1
            i += 1  # closing fence
            yield ("code", body)
            continue

        # Pipe table — header row | --- | --- |
        if "|" in line and i + 1 < n and re.match(r"^\s*\|?\s*:?-+", lines[i + 1]):
            header = [c.strip() for c in line.strip().strip("|").split("|")]
            i += 2  # skip header + separator
            rows = []
            while i < n and "|" in lines[i] and lines[i].strip():
                row = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                rows.append(row)
                i += 1
            yield ("table", header, rows)
            continue

        # Bulleted list
        if re.match(r"^\s*[-\*]\s+", line):
            while i < n and re.match(r"^\s*[-\*]\s+", lines[i]):
                bullet_text = re.sub(r"^\s*[-\*]\s+", "", lines[i])
                yield ("bullet", bullet_text)
                i += 1
            continue

        # Numbered list \u2014 keep as bullets for simplicity
        if re.match(r"^\s*\d+\.\s+", line):
            while i < n and re.match(r"^\s*\d+\.\s+", lines[i]):
                bullet_text = re.sub(r"^\s*\d+\.\s+", "", lines[i])
                yield ("bullet", bullet_text)
                i += 1
            continue

        # Paragraph: collect until blank
        para_lines = [line]
        i += 1
        while i < n and lines[i].strip() and not (
            lines[i].lstrip().startswith("#")
            or lines[i].strip().startswith("```")
            or re.match(r"^\s*[-\*]\s+", lines[i])
            or re.match(r"^\s*\d+\.\s+", lines[i])
            or "|" in lines[i] and i + 1 < n and re.match(r"^\s*\|?\s*:?-+", lines[i + 1])
            or re.fullmatch(r"-{3,}", lines[i].strip())
        ):
            para_lines.append(lines[i])
            i += 1
        yield ("para", " ".join(p.strip() for p in para_lines))


# ── Document building ────────────────────────────────────────────────────────
def build():
    md = SRC.read_text(encoding="utf-8")
    doc = Document()

    # Page setup: A4-friendly margins
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)

    # Default font
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # Cover-style title block (manually placed)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_run = title.add_run("ELSA Depression Prediction")
    title_run.bold = True
    title_run.font.size = Pt(28)
    title_run.font.color.rgb = NAVY
    title_run.font.name = "Calibri"

    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(2)
    sub_run = sub.add_run(
        "Early prediction from 4 years out using ELSA longitudinal data"
    )
    sub_run.italic = True
    sub_run.font.size = Pt(13)
    sub_run.font.color.rgb = TEAL
    sub_run.font.name = "Calibri"

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(2)
    meta_run = meta.add_run(
        "AI and Healthcare \u2014 University of Surrey, MSc Artificial Intelligence (2025/26)"
    )
    meta_run.font.size = Pt(11)
    meta_run.font.color.rgb = GREY

    team = doc.add_paragraph()
    team.paragraph_format.space_after = Pt(2)
    team_run = team.add_run(
        "Akeeb Lawel  \u2022  Fiyin Akano  \u2022  Zannat Chowdhury Sagar  \u2022  "
        "Giridhar Nampally  \u2022  Pushkar Jadav  \u2022  Poorna Golla"
    )
    team_run.font.size = Pt(10.5)
    team_run.font.color.rgb = GREY

    _hr(doc)

    # Process Markdown
    skip_first_h1 = True
    for tok in parse_markdown(md):
        kind = tok[0]
        if kind == "heading":
            level, text = tok[1], tok[2]
            if skip_first_h1 and level == 1:
                skip_first_h1 = False
                continue
            _heading(doc, text, level)
        elif kind == "para":
            _paragraph(doc, tok[1])
        elif kind == "bullet":
            _bullet(doc, tok[1])
        elif kind == "code":
            _code_block(doc, tok[1])
        elif kind == "table":
            _table(doc, tok[1], tok[2])
        elif kind == "hr":
            _hr(doc)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"Saved DOCX: {OUT}")
    print(f"Size: {OUT.stat().st_size // 1024} KB")


if __name__ == "__main__":
    build()
