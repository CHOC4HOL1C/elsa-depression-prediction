"""
Append Meeting 7 (final presentation rehearsal) to the existing
docs/Meeting_Minutes.docx, preserving the user's manual arrangement.

Run:    python3 docs/append_meeting_7.py
Output: docs/Meeting_Minutes.docx (modified in place)

This script is idempotent — it checks for an existing 'MEETING MINUTES [7]'
heading and exits without changes if one is already present, so re-running
is safe.
"""
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor

SRC = Path(__file__).resolve().parent / "Meeting_Minutes.docx"
NAVY = RGBColor(0x1F, 0x2D, 0x5C)


MEETING_7 = dict(
    number=7,
    date="01/05/2026",
    time="11am",
    attendance=[
        "Lawal, Akeeb",
        "Zannat Chowdhury",
        "Akano Fiyinfoluwa",
        "Golla, Poorna Chandra C",
        "Jadhav, Pushkar Vijay V",
        "Nampally, Giridhar Kumar",
    ],
    goal="Final presentation rehearsal and submission sign-off",
    discussion=[
        "Final team rehearsal in advance of the Week 11 in-class presentation.",
        "Akeeb walked through the background, ELSA dataset, and wave-selection sections of the deck (slides 2\u20134), explaining the audit findings.",
        "Fiyin presented the modelling pipeline, experimental design, and headline results (slides 5\u20138), including the SHAP interpretation.",
        "Zannat presented the leakage sensitivity finding and the limitations / future directions sections (slides 9\u201312).",
        "Team discussed slide ordering, timing, and the story arc; agreed to lead with the leakage sensitivity result on slide 10 as the central scientific contribution.",
        "Conducted a full timed run-through; landed at 9 minutes 40 seconds, comfortable for the 10-minute slot.",
        "Discussed Q&A preparation: reviewed the QA_FLASHCARDS pack and assigned topic ownership for likely questions.",
        "Reviewed all final submission artefacts: the run-once Colab notebook, the report (DOCX + Markdown), the team contributions document, and the meeting minutes.",
        "Confirmed submission package contents and SurreyLearn upload plan.",
    ],
    actions=[
        "Akeeb: Open the talk by presenting sections 1\u20133 (background, ELSA dataset, wave selection).",
        "Fiyin: Present sections 4\u20137 (pipeline, methodology, headline results, SHAP).",
        "Zannat: Present sections 8\u201310 (leakage sensitivity, limitations, conclusion).",
        "Pushkar: Timekeeper during the live presentation; prompt speakers if approaching the 10-minute limit.",
        "Giridhar: Backup Q&A for methodology questions; have the QA_FLASHCARDS open during delivery.",
        "Poorna: Backup Q&A for background and dataset questions; circulate the final meeting minutes to the team.",
        "Everyone: Familiarise with QA_FLASHCARDS.md before delivery day; one more dry-run scheduled informally.",
    ],
    next_meeting="Live presentation \u2014 Week 11 class",
)


def _add_para(doc, text, bold=False, size=11, color=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    if color:
        run.font.color.rgb = color
    return p


def _add_blank(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    return p


def append_meeting(doc, m):
    _add_blank(doc)
    _add_para(doc, f"MEETING MINUTES [{m['number']}]",
              bold=True, size=14, color=NAVY)
    _add_blank(doc)

    _add_para(doc, "Date & Time:", bold=True)
    _add_para(doc, f"DAY: {m['date']}")
    _add_para(doc, f"TIME: {m['time']}")
    _add_blank(doc)

    _add_para(doc, "Attendance:", bold=True)
    for a in m["attendance"]:
        _add_para(doc, a)
    _add_blank(doc)

    _add_para(doc, "Today\u2019s Goal:", bold=True)
    _add_para(doc, m["goal"])
    _add_blank(doc)

    _add_para(doc, "Discussion Points:", bold=True)
    for d in m["discussion"]:
        _add_para(doc, d)
    _add_blank(doc)

    _add_para(doc, "Action Items (Who is doing what?):", bold=True)
    for a in m["actions"]:
        _add_para(doc, a)
    _add_blank(doc)

    _add_para(doc, f"Next Meeting: {m['next_meeting']}", bold=True)


def main():
    if not SRC.exists():
        raise FileNotFoundError(f"Source minutes not found: {SRC}")

    doc = Document(str(SRC))

    # Idempotency guard: do not append if Meeting 7 already exists
    for p in doc.paragraphs:
        if "MEETING MINUTES [7]" in p.text.upper().replace(" ", " "):
            print("Meeting 7 already exists \u2014 nothing to do.")
            return

    append_meeting(doc, MEETING_7)
    doc.save(str(SRC))
    print(f"Appended Meeting 7 to: {SRC}")
    print(f"Size: {SRC.stat().st_size // 1024} KB")


if __name__ == "__main__":
    main()
