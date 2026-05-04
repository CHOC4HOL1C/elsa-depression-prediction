"""
Build the project coordination log document from scratch.

Replaces the previous build_meeting_minutes.py + append_meeting_7.py pair.

This document records all forms of team coordination over the project, not
only synchronous meetings. Each entry is labelled by activity type:

  - MEETING (Live)    — full-team or majority-team synchronous calls
  - WORKING SESSION   — focused sub-team work between two or more members
  - PR REVIEW (Async) — coordination via GitHub pull requests
  - ASYNC SYNC        — short asynchronous coordination (group chat, threads)

Output: docs/Meeting_Minutes.docx
        (Filename retained for the SurreyLearn coursework brief, which asks
         specifically for "meeting minutes". The document title and per-entry
         labelling reflect the broader scope.)
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor

OUT = Path(__file__).resolve().parent / "Meeting_Minutes.docx"

NAVY = RGBColor(0x1F, 0x2D, 0x5C)
INK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x55, 0x55, 0x55)


# ── Coordination entries (chronological) ────────────────────────────────────
ENTRIES = [
    dict(
        kind="MEETING (Live)",
        number=1,
        date="26/02/2026",
        time="6pm",
        participants=[
            "Lawal, Akeeb",
            "Zannat Chowdhury",
            "Akano Fiyinfoluwa",
            "Golla, Poorna Chandra C",
            "Jadhav, Pushkar Vijay V",
            "Nampally, Giridhar Kumar",
        ],
        goal="Project kickoff and introductions",
        discussion=[
            "Introduction of group members; exchange of pleasantries.",
            "Discussion of next steps and how the team would coordinate going forward.",
            "Decision to schedule online meetings as the default coordination format.",
            "Selection of WhatsApp group chat as the asynchronous communication channel.",
        ],
        actions=[
            "Everyone: Read up on ELSA.",
            "Akeeb: Register on the UK Data Service to obtain access to the ELSA dataset.",
            "Fiyin: Create a centralised tracking document for project objectives and tasks.",
        ],
        next_meeting="TBD",
    ),
    dict(
        kind="MEETING (Live)",
        number=2,
        date="08/04/2026",
        time="10am",
        participants=[
            "Lawal, Akeeb",
            "Zannat Chowdhury",
            "Akano Fiyinfoluwa",
            "Jadhav, Pushkar Vijay V",
            "Nampally, Giridhar Kumar",
        ],
        goal="Roadmap discussion and task allocation",
        discussion=[
            "Discussion of the next steps of the project and how it would be delivered.",
            "Allocation of initial tasks across the team.",
        ],
        actions=[
            "Everyone: Get hands-on with the ELSA dataset.",
            "Akeeb: Lead data preprocessing and audit.",
            "Fiyin: Lead model creation.",
            "Zannat: Work on data loading.",
            "Pushkar: Investigate feature selection.",
        ],
        next_meeting="10/04/2026",
    ),
    dict(
        kind="MEETING (Live)",
        number=3,
        date="10/04/2026",
        time="2pm",
        participants=[
            "Lawal, Akeeb",
            "Zannat Chowdhury",
            "Akano Fiyinfoluwa",
            "Golla, Poorna Chandra C",
            "Jadhav, Pushkar Vijay V",
            "Nampally, Giridhar Kumar",
        ],
        goal="Wave selection and methodology framing",
        discussion=[
            "Akeeb presented findings from his ELSA core wave audit (Waves 3–11).",
            "Reviewed CES-D quality, participant overlap, and feature drift across waves.",
            "Confirmed Waves 6, 7, and 8 form the strongest consecutive longitudinal block "
            "(>90% participant overlap, 93–97% valid CES-D responses).",
            "Locked the modelling design: predict CES-D ≥ 3 at Wave 8 from Wave 6 and "
            "Wave 7 features (binary classification).",
            "Agreed on Logistic Regression as the interpretable baseline, with Random "
            "Forest for non-linear comparison.",
            "Noted that the IFS-derived variable files contain the validated cesd_sc; "
            "agreed to use these rather than reconstructing the score from raw items.",
        ],
        actions=[
            "Akeeb: Finalise the curated feature list from the audit "
            "(target ~30 predictors per wave).",
            "Fiyin: Begin building the W6+W7 modelling pipeline with CES-D as the primary outcome.",
            "Zannat: Validate preprocessing assumptions on a single-wave baseline.",
            "Pushkar: Investigate feature selection methods.",
            "Giridhar & Poorna: Read up on classification evaluation metrics for "
            "imbalanced classes (ROC-AUC, F1, PR curves).",
        ],
        next_meeting="14/04/2026",
    ),
    dict(
        kind="PR REVIEW (Async)",
        number=1,
        date="11/04/2026 – 13/04/2026",
        time="Asynchronous via GitHub",
        participants=[
            "Akeeb (PR author)",
            "Fiyin (reviewer)",
            "Zannat (reviewer)",
        ],
        goal="Review and approve the data audit notebook before the modelling work began",
        discussion=[
            "Akeeb opened a pull request containing the full audit notebook and the "
            "curated feature list (modelling_features_final_curated.csv).",
            "Reviewers verified reproducibility: the notebook runs end-to-end on the "
            "team-shared dataset folder.",
            "Confirmed the wave-selection rationale and the 30-feature funnel logic.",
            "Minor documentation comments were raised and addressed by Akeeb in "
            "follow-up commits.",
            "PR approved and merged; wave selection was officially locked for "
            "downstream modelling.",
        ],
        actions=[
            "Akeeb: Merge the audit branch and tag the latest commit as the audit baseline.",
            "Fiyin & Zannat: Pull the merged audit branch into modelling work.",
        ],
        next_meeting="—",
    ),
    dict(
        kind="ASYNC SYNC",
        number=1,
        date="14/04/2026",
        time="Group chat thread",
        participants=[
            "Akano Fiyinfoluwa",
            "Lawal, Akeeb",
        ],
        goal="CES-D scoring fix decision and modelling progress check",
        discussion=[
            "Fiyin reported on the W6+W7 pipeline build over the GC: data loading, "
            "merging, preprocessing, and Logistic Regression training all functional.",
            "Fiyin flagged that an earlier exploratory CES-D reconstruction (from raw "
            "PSced items, 0–3 scale) produced an apparent prevalence of ~74%, which is "
            "clinically implausible.",
            "Investigation traced this to a scale mismatch: ELSA's CES-D items are "
            "binary (1 = yes / 2 = no after Stata recoding), not 0–3.",
            "Resolution: switch to the IFS-validated cesd_sc variable (0–8 range). "
            "Prevalence corrects to ~19%, consistent with published ELSA estimates.",
            "Discussed adding a leakage sensitivity analysis (with vs without prior "
            "CES-D as predictors) to test whether the model is just \u201cdepressed people "
            "stay depressed\u201d.",
            "Agreed to add SHAP analysis for model interpretability.",
        ],
        actions=[
            "Fiyin: Implement the CES-D fix; add the leakage sensitivity test; "
            "integrate SHAP explanations.",
            "Zannat: Begin design of the enhanced pipeline (additional models, "
            "hyperparameter tuning, calibration).",
            "Akeeb: Document the CES-D scoring lesson in CONTEXT.md.",
        ],
        next_meeting="16/04/2026 (Fiyin & Zannat working session)",
    ),
    dict(
        kind="WORKING SESSION",
        number=1,
        date="16/04/2026",
        time="3pm",
        participants=[
            "Akano Fiyinfoluwa",
            "Zannat Chowdhury",
        ],
        goal="Modelling sub-team alignment on the enhanced pipeline and submission integration",
        discussion=[
            "Working session between the two members owning the modelling pipeline.",
            "Zannat presented the enhanced longitudinal pipeline: three prediction arms "
            "(W6, W7, W6+W7), five feature configurations per arm, four models "
            "(LR, RF, XGBoost, LightGBM), hyperparameter tuning via RandomizedSearchCV, "
            "and ensemble averaging.",
            "Reviewed the SHAP \u201ctwo-act\u201d narrative: one analysis with the full feature "
            "set (showing prior CES-D dominance) and one without CES-D (revealing "
            "self-rated health, mobility, and financial difficulty as the top "
            "independent predictors).",
            "Agreed on the structure of the final submission notebook: single file, "
            "run-once on Colab, with auto-installation of dependencies and one "
            "user-set ELSA_PATH variable.",
            "Reviewed the leakage sensitivity numbers: AUC 0.85 with prior CES-D "
            "→ 0.77 without. Agreed this is the central scientific finding.",
            "Discussed presentation slide ordering and story arc.",
        ],
        actions=[
            "Zannat: Finalise the submission notebook; add inline result interpretation "
            "between cells.",
            "Fiyin: Draft the project report and presentation deck; integrate "
            "Akeeb's audit findings into the background section.",
            "Both: Review the final notebook end-to-end before submission.",
        ],
        next_meeting="—",
    ),
    dict(
        kind="PR REVIEW (Async)",
        number=2,
        date="17/04/2026 – 28/04/2026",
        time="Asynchronous via GitHub (multiple PR cycles)",
        participants=[
            "Akano Fiyinfoluwa (PR author for modelling pipeline)",
            "Zannat Chowdhury (PR author for enhanced pipeline and submission notebook)",
            "GitHub Copilot (automated reviewer)",
            "Cross-reviewers: Fiyin and Zannat reviewed each other's PRs",
        ],
        goal="Iterative review of the modelling, enhanced pipeline, and submission notebook PRs",
        discussion=[
            "PRs #2 through #11 covered the W6+W7 modelling pipeline, the enhanced "
            "pipeline (adding XGBoost, LightGBM, tuning, calibration, SHAP, domain "
            "ablation), and the final submission notebook.",
            "Eleven rounds of GitHub Copilot feedback were addressed, including: "
            "removing personal paths from config.py, robust REPO_ROOT discovery, "
            "training-only computation of the 40% missingness drop threshold "
            "(eliminating test-set leakage in feature selection), corrected axis "
            "labelling on the standardised LR coefficient plot, and the wave-symmetry "
            "rule for missingness drops.",
            "Zannat's review of Fiyin's PR #2 identified the binary F1 framing, "
            "flagged the missing XGBoost comparison, and approved the merge with "
            "substantive written commentary.",
            "Some Copilot suggestions (CV-fold-internal missingness filtering, "
            "lazy data path scanning, partial-sum imputation for functional_burden) "
            "were acknowledged as documented limitations rather than fixes, where "
            "the practical impact on results was negligible.",
        ],
        actions=[
            "Fiyin: Address Copilot review comments on the modelling pipeline PR; "
            "merge to dev.",
            "Zannat: Address Copilot review comments on the enhanced pipeline PR; "
            "merge to dev.",
            "Zannat: Open and merge the final submission notebook PR.",
        ],
        next_meeting="—",
    ),
    dict(
        kind="MEETING (Live)",
        number=4,
        date="01/05/2026",
        time="11am",
        participants=[
            "Lawal, Akeeb",
            "Zannat Chowdhury",
            "Akano Fiyinfoluwa",
            "Golla, Poorna Chandra C",
            "Jadhav, Pushkar Vijay V",
            "Nampally, Giridhar Kumar",
        ],
        goal="Final presentation rehearsal and submission sign-off",
        discussion=[
            "Team reviewed slide ordering, timing, and the story arc.",
            "Conducted a full timed run-through; landed at 9 minutes 40 seconds — "
            "comfortable for the 10-minute slot.",
            "Reviewed the QA_FLASHCARDS pack and assigned topic ownership for likely "
            "questions.",
            "Reviewed all final submission artefacts: the run-once Colab notebook, "
            "the report (DOCX + Markdown), the team contributions document, and this "
            "coordination log.",
            "Confirmed submission package contents and the SurreyLearn upload plan.",
        ],
        actions=[
            "Akeeb: Open the talk by presenting sections 1–3 (background, ELSA "
            "dataset, wave selection).",
            "Fiyin: Present sections 4–7 (pipeline, methodology, headline results, "
            "SHAP).",
            "Zannat: Present sections 8–10 (leakage sensitivity, limitations, "
            "conclusion).",
            "Pushkar: Timekeeper during the live presentation; prompt speakers if "
            "they approach the 10-minute limit.",
            "Giridhar: Backup Q&A for methodology questions; have the QA_FLASHCARDS "
            "open during delivery.",
            "Poorna: Backup Q&A for background and dataset questions; circulate the "
            "final coordination log to the team.",
            "Everyone: Familiarise with QA_FLASHCARDS.md before delivery day.",
        ],
        next_meeting="Live presentation — Week 11 class",
    ),
]


# ── Document helpers ────────────────────────────────────────────────────────
def _add_para(doc, text, bold=False, size=11, color=None, italic=False, align=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    if color is not None:
        run.font.color.rgb = color
    return p


def _add_blank(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    return p


def add_header(doc):
    _add_para(doc, "University of Surrey", bold=True, size=14, color=NAVY)
    _add_para(doc, "AI & Healthcare (EEEM069)", bold=True, size=12, color=NAVY)
    _add_para(doc, "Group 4 — \u201CHealth is Wealth\u201D", bold=True, size=12, color=NAVY)
    _add_para(
        doc,
        "Members: Lawal Akeeb, Zannat Chowdhury, Akano Fiyinfoluwa, "
        "Golla Poorna Chandra, Jadhav Pushkar Vijay, Nampally Giridhar Kumar",
        size=10,
        color=GREY,
    )
    _add_blank(doc)
    _add_para(doc, "Project Coordination Log", bold=True, size=16, color=NAVY)
    _add_para(
        doc,
        "This document records the team's coordination throughout the project. "
        "It includes synchronous group meetings, asynchronous coordination via "
        "GitHub pull-request reviews, sub-team working sessions, and group-chat "
        "discussions. Entries are listed in chronological order and labelled by "
        "activity type so that the reader can distinguish a live meeting from an "
        "asynchronous review.",
        italic=True,
        size=10,
        color=GREY,
    )
    _add_blank(doc)
    _add_para(doc, "Activity types in this log:", bold=True, size=11)
    _add_para(doc, "  • MEETING (Live)    — full-team or majority-team synchronous call.", size=10)
    _add_para(doc, "  • PR REVIEW (Async) — coordination via GitHub pull request.", size=10)
    _add_para(doc, "  • WORKING SESSION   — focused sub-team work between two or more members.", size=10)
    _add_para(doc, "  • ASYNC SYNC        — short asynchronous coordination (group chat).", size=10)
    _add_blank(doc)


def add_entry(doc, entry):
    _add_blank(doc)
    title = f"{entry['kind']}  [{entry['number']}]"
    _add_para(doc, title, bold=True, size=14, color=NAVY)
    _add_blank(doc)

    _add_para(doc, "Date & Time:", bold=True)
    _add_para(doc, f"DATE: {entry['date']}")
    _add_para(doc, f"TIME / FORMAT: {entry['time']}")
    _add_blank(doc)

    _add_para(doc, "Participants:", bold=True)
    for a in entry["participants"]:
        _add_para(doc, f"  • {a}")
    _add_blank(doc)

    _add_para(doc, "Goal:", bold=True)
    _add_para(doc, entry["goal"])
    _add_blank(doc)

    _add_para(doc, "What was discussed / done:", bold=True)
    for d in entry["discussion"]:
        _add_para(doc, f"  • {d}")
    _add_blank(doc)

    _add_para(doc, "Action items (who is doing what):", bold=True)
    for a in entry["actions"]:
        _add_para(doc, f"  • {a}")
    _add_blank(doc)

    if entry.get("next_meeting") and entry["next_meeting"] != "—":
        _add_para(doc, f"Next coordination point: {entry['next_meeting']}", bold=True)


def main():
    doc = Document()

    # Document margins
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Pt(60)
        section.left_margin = section.right_margin = Pt(60)

    add_header(doc)

    for entry in ENTRIES:
        add_entry(doc, entry)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"Saved: {OUT}")
    print(f"Size: {OUT.stat().st_size // 1024} KB")
    print(f"Total entries: {len(ENTRIES)}")
    kinds = {}
    for e in ENTRIES:
        k = e["kind"]
        kinds[k] = kinds.get(k, 0) + 1
    print("Breakdown:")
    for k, n in kinds.items():
        print(f"  {k}: {n}")


if __name__ == "__main__":
    main()
